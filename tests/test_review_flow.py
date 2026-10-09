import csv
import json
import tempfile
import unittest
from pathlib import Path
from primate_ai.behavior_review import (
    ReviewStore, CameraTimeIndex, camera_sources, codebook_path,
    save_codebook,read_codebook,export_workbook,export_batch_workbook,stage_windows
)


class ReviewWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)

    def test_review_persists_and_is_idempotent(self):
        with (self.root/'events.csv').open('w',encoding='utf-8',newline='') as f:
            w=csv.DictWriter(f,fieldnames=['event_id','start_s','end_s','peak_s','candidate_label','support_cameras'])
            w.writeheader()
            w.writerow(dict(event_id='1',start_s=.1,end_s=.2,peak_s=.15,candidate_label='twitch-like',support_cameras='cam01,cam02'))
            w.writerow(dict(event_id='2',start_s=1.1,end_s=1.2,peak_s=1.15,candidate_label='tremor-like',support_cameras='cam03'))
        store=ReviewStore(self.root)
        self.assertEqual(len(store.events),2)
        store.update(0,review='confirmed',label='抽搐样短促运动',note='复核确认')
        created=store.add('静止/休息',3.,7.,review='confirmed')
        self.assertEqual(created,'M0001')
        again=ReviewStore(self.root)
        self.assertEqual(len(again.events),3)
        self.assertEqual(again.events[0]['review'],'confirmed')
        self.assertEqual(again.events[0]['note'],'复核确认')
        self.assertTrue(any(x['event_id']=='M0001' for x in again.events))

    def test_invalid_duration_is_clamped(self):
        s=ReviewStore(self.root)
        s.add('test',2,1)
        self.assertEqual(s.events[0]['start_s'],s.events[0]['end_s'])
        self.assertTrue((self.root/'review_annotations.csv').exists())

    def test_camera_timestamp_nearest_frame_and_fallback(self):
        with (self.root/'timestamps_cam01_front.csv').open('w',encoding='utf-8',newline='') as f:
            w=csv.DictWriter(f,fieldnames=['frame_index','session_time_s']);w.writeheader()
            for n,t in enumerate([.06,.10,.14,.18]):w.writerow(dict(frame_index=n,session_time_s=t))
        idx=CameraTimeIndex(self.root,'cam01',30)
        self.assertEqual(idx.frame_at(.15),2)
        self.assertAlmostEqual(idx.time_at(2),.14,places=3)
        fallback=CameraTimeIndex(self.root,'cam02',30)
        self.assertEqual(fallback.frame_at(1.),30)

    def test_camera_manifest_one_to_three(self):
        details=self.root/'details';details.mkdir()
        cams=[]
        for n in range(1,4):
            video=self.root/('cam%02d_front.avi'%n)
            video.write_bytes(b'placeholder')
            cams.append(dict(camera_id='cam%02d'%n,view='front',video_path=str(video)))
        (details/'analysis_manifest.json').write_text(json.dumps(dict(session_dir=str(self.root),camera_inputs=cams)),encoding='utf-8')
        self.assertEqual(len(camera_sources(self.root)),3)

    def test_custom_codebook(self):
        save_codebook(self.root,[dict(label='头部活动',kind='state'),dict(label='抽动',kind='point')])
        self.assertTrue(codebook_path(self.root).exists())
        self.assertEqual(read_codebook(self.root)[1]['kind'],'point')

    def test_stage_timeline_and_batch_compare(self):
        details=self.root/'details';details.mkdir()
        (details/'analysis_manifest.json').write_text(json.dumps({'session_dir':str(self.root),'camera_count':2}),encoding='utf-8')
        meta=self.root/'_metadata';meta.mkdir()
        (meta/'analysis_windows.csv').write_text('window_type,condition,start_s,end_s\nBaseline,,0,30\nExposure,EMF40,30,40\nRecovery,,40,70\n',encoding='utf-8')
        self.assertEqual(len(stage_windows(self.root)),3)
        (details/'stage_summary.csv').write_text('camera_id,stage,condition,covered_s\ncam01,Baseline,,30\ncam02,Exposure,EMF40,10\n',encoding='utf-8')
        (self.root/'summary.csv').write_text('camera_id,metric,value\ncam01,activity,0.5\n',encoding='utf-8')
        ReviewStore(self.root).add('抽搐样短促运动',36,36.2,review='confirmed')
        dst=self.root/'compare.xlsx'
        export_batch_workbook([self.root],dst)
        from openpyxl import load_workbook
        book=load_workbook(dst,read_only=True)
        self.assertEqual(book['实验列表']['D2'].value,1)
        self.assertEqual(book['按阶段各视角统计'].max_row,3)
        book.close()

    def test_excel_workbook(self):
        (self.root/'details').mkdir()
        (self.root/'summary.csv').write_text('camera_id,metric,value\ncam01,activity,0.7\n',encoding='utf-8')
        (self.root/'details'/'stage_summary.csv').write_text('camera_id,stage,mean_activity_proxy\ncam01,Baseline,0.5\n',encoding='utf-8')
        store=ReviewStore(self.root)
        store.add('twitch',1.0,1.2,review='confirmed')
        xlsx=export_workbook(self.root)
        self.assertTrue(xlsx.exists())
        from openpyxl import load_workbook
        book=load_workbook(xlsx,read_only=True)
        self.assertEqual(book.sheetnames,['实验信息','阶段统计','核心指标','事件与人工复核','人工确认统计'])
        self.assertEqual(book['人工确认统计']['B2'].value,1)
        book.close()

if __name__=='__main__':unittest.main()
